from io import BytesIO
import os

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt
from fpdf import FPDF

from backend.utils.text import sanitize_text


# ============================================================
# TXT EXPORT
# ============================================================

def make_txt(text: str) -> bytes:
    return sanitize_text(text).encode("utf-8")


# ============================================================
# DOCX EXPORT
# ============================================================

def make_docx(
    text: str,
    document_type: str = "Legal Document",
    terms: str = "",
    brand_name: str = "LegalEase",
) -> bytes:

    document = Document()
    section = document.sections[0]

    section.top_margin = Inches(0.7)
    section.bottom_margin = Inches(0.7)
    section.left_margin = Inches(0.8)
    section.right_margin = Inches(0.8)

    styles = document.styles
    styles["Normal"].font.name = "Arial"
    styles["Normal"].font.size = Pt(10.5)

    # Header
    header = section.header
    p = header.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER

    r = p.add_run(brand_name)
    r.bold = True
    r.font.size = Pt(14)

    # Title
    title = document.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER

    r = title.add_run(document_type.upper())
    r.bold = True
    r.font.size = Pt(16)

    document.add_paragraph()

    cleaned = sanitize_text(text)

    for block in cleaned.split("\n\n"):

        block = block.strip()

        if not block:
            continue

        paragraph = document.add_paragraph()

        for index, line in enumerate(block.splitlines()):

            line = line.strip()

            if not line:
                continue

            if index > 0:
                paragraph.add_run("\n")

            r = paragraph.add_run(line)

            if line.isupper():
                r.bold = True

    # Footer
    footer = section.footer
    p = footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER

    r = p.add_run(
        f"{brand_name} | AI-assisted draft | Not legal advice"
    )
    r.font.size = Pt(8)

    output = BytesIO()
    document.save(output)

    return output.getvalue()


# ============================================================
# PDF CLASS
# ============================================================

class LegalPDF(FPDF):

    def __init__(
        self,
        brand_name="LegalEase",
        document_type="Legal Document",
        logo_path=None,
    ):
        super().__init__()

        self.brand_name = brand_name
        self.document_type = document_type
        self.logo_path = logo_path

        # Margins
        # Top margin is larger because logo is displayed there
        self.set_margins(
            18,
            32,
            18
        )

        # Space for footer
        self.set_auto_page_break(
            auto=True,
            margin=20
        )

    # --------------------------------------------------------
    # PDF HEADER
    # --------------------------------------------------------

    def header(self):

        # Display logo at top center
        if self.logo_path and os.path.exists(self.logo_path):

            logo_width = 35

            x_position = (
                self.w - logo_width
            ) / 2

            self.image(
                self.logo_path,
                x=x_position,
                y=8,
                w=logo_width
            )

            # Move cursor below logo
            self.set_y(28)

        else:

            # If logo is not found
            self.set_y(10)

        # Brand name
        self.set_font(
            "Helvetica",
            "B",
            13
        )

        self.cell(
            0,
            8,
            self.brand_name,
            align="C"
        )

        self.ln(8)

        # Document title
        self.set_font(
            "Helvetica",
            "B",
            14
        )

        self.multi_cell(
            0,
            8,
            self.document_type.upper(),
            align="C"
        )

        self.ln(6)

    # --------------------------------------------------------
    # PDF FOOTER
    # --------------------------------------------------------

    def footer(self):

        self.set_y(-15)

        self.set_font(
            "Helvetica",
            "",
            8
        )

        self.cell(
            0,
            10,
            f"{self.brand_name} | AI-assisted draft | Not legal advice",
            align="C"
        )


# ============================================================
# SAFE PDF TEXT
# ============================================================

def safe_pdf_text(text: str) -> str:

    replacements = {
        "\u2018": "'",
        "\u2019": "'",
        "\u201c": '"',
        "\u201d": '"',
        "\u2013": "-",
        "\u2014": "-",
        "\u2022": "-",
        "\u00a0": " ",
        "\u2026": "...",
        "\u00a9": "(c)",
        "\u00ae": "(R)",
        "\u2122": "(TM)",
    }

    text = str(text)

    for old, new in replacements.items():
        text = text.replace(old, new)

    return text.encode(
        "latin-1",
        "replace"
    ).decode("latin-1")


# ============================================================
# PDF EXPORT
# ============================================================

def make_pdf(
    text: str,
    document_type: str = "Legal Document",
    terms: str = "",
    brand_name: str = "LegalEase",
) -> bytes:

    # --------------------------------------------------------
    # LOGO LOCATION
    # --------------------------------------------------------

    logo_path = os.path.join(
        os.path.dirname(__file__),
        "..",
        "assets",
        "logo.png"
    )

    logo_path = os.path.abspath(logo_path)

    # --------------------------------------------------------
    # CREATE PDF
    # --------------------------------------------------------

    pdf = LegalPDF(
        brand_name=brand_name,
        document_type=document_type,
        logo_path=logo_path
    )

    # PDF metadata
    pdf.set_title(
        f"{document_type} - {brand_name}"
    )

    pdf.set_author(
        brand_name
    )

    # Add first page
    pdf.add_page()

    # Normal text
    pdf.set_font(
        "Helvetica",
        "",
        10
    )

    # Clean text
    cleaned = sanitize_text(text)
    cleaned = safe_pdf_text(cleaned)

    paragraphs = cleaned.split("\n\n")

    # Write paragraphs
    for paragraph in paragraphs:

        paragraph = paragraph.strip()

        if not paragraph:
            continue

        pdf.multi_cell(
            0,
            6,
            paragraph
        )

        pdf.ln(3)

    # Return PDF bytes
    return bytes(pdf.output())